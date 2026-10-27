"""Informe Excel listo para el equipo contable."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = "1F4E78"
GRIS_TEXTO = "8C8C8C"
FONDO_CONCILIADO = "EDEDED"
FONDO_PENDIENTE = "FFF2CC"
ROJO_ALERTA = "F8CBAD"
MONEDA = '#,##0.00;[Red]-#,##0.00'
FECHA = "DD/MM/YYYY"

COLUMNAS_DETALLE = {
    "fecha": "Fecha", "origen": "Origen", "concepto": "Concepto", "ref": "Referencia",
    "categoria": "Cuenta contable", "nombre": "Contraparte", "ingreso": "Ingreso",
    "egreso": "Egreso", "conciliado": "Conciliado", "estado": "Estado", "nivel": "Nivel",
    "match": "Match", "dias_diferencia": "Días dif.", "usuario": "Usuario",
    "fecha_registro": "Fecha registro",
}


def _resumen(resultados) -> pd.DataFrame:
    filas = []
    for r in resultados:
        fila = {"Cuenta": r.cuenta, "Movimientos": r.total, "Conciliados": r.conciliados,
                "% automático": r.tasa, "Pendientes": r.total - r.conciliados}
        fila.update(r.conciliados_por_nivel)
        filas.append(fila)
    return pd.DataFrame(filas).fillna(0)


def _por_cuenta_contable(resultados) -> pd.DataFrame:
    partes = []
    for r in resultados:
        g = (r.movimientos.groupby(["categoria", "origen"])[["ingreso", "egreso"]].sum()
             .unstack(fill_value=0))
        g.columns = [f"{a.capitalize()} {b}" for a, b in g.columns]
        for col in ("Ingreso Banco", "Ingreso Contabilidad", "Egreso Banco", "Egreso Contabilidad"):
            if col not in g:
                g[col] = 0.0
        g["Dif. ingreso"] = g["Ingreso Banco"] - g["Ingreso Contabilidad"]
        g["Dif. egreso"] = g["Egreso Banco"] - g["Egreso Contabilidad"]
        g.insert(0, "Banco", r.cuenta)
        partes.append(g.reset_index().rename(columns={"categoria": "Cuenta contable"}))
    return pd.concat(partes, ignore_index=True)[
        ["Banco", "Cuenta contable", "Ingreso Banco", "Ingreso Contabilidad", "Dif. ingreso",
         "Egreso Banco", "Egreso Contabilidad", "Dif. egreso"]]


def _formatear(ws, columnas_moneda=(), columnas_fecha=(), anchos=None):
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    encabezados = [c.value for c in ws[1]]
    for i, nombre in enumerate(encabezados, 1):
        letra = get_column_letter(i)
        ancho = (anchos or {}).get(nombre, max(12, min(60, len(str(nombre)) + 4)))
        ws.column_dimensions[letra].width = ancho
        fmt = MONEDA if nombre in columnas_moneda else FECHA if nombre in columnas_fecha else None
        if fmt:
            for celda in ws[letra][1:]:
                celda.number_format = fmt


def exportar_excel(resultados, pendientes: pd.DataFrame, ruta: str | Path,
                   control: pd.DataFrame | None = None) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)

    resumen = _resumen(resultados)
    cols_pend = ["cuenta", "fecha", "origen", "concepto", "ingreso", "egreso", "tipo", "sugerencia", "estado"]
    hoja_pend = pendientes[cols_pend].rename(columns={
        "cuenta": "Banco", "fecha": "Fecha", "origen": "Origen", "concepto": "Concepto",
        "ingreso": "Ingreso", "egreso": "Egreso", "tipo": "Tipo", "sugerencia": "Sugerencia",
        "estado": "Estado"}).sort_values(["Banco", "Fecha"])

    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        resumen.to_excel(xw, sheet_name="Resumen", index=False)
        hoja_pend.to_excel(xw, sheet_name="Pendientes", index=False)
        if control is not None and not control.empty:
            control.to_excel(xw, sheet_name="Control de matches", index=False)
        _por_cuenta_contable(resultados).to_excel(xw, sheet_name="Por cuenta contable", index=False)
        for r in resultados:
            det = (r.movimientos.sort_values(["origen", "fecha"])[list(COLUMNAS_DETALLE)]
                   .rename(columns=COLUMNAS_DETALLE))
            det.to_excel(xw, sheet_name=r.cuenta[:31], index=False)

    wb = load_workbook(ruta)
    _formatear(wb["Resumen"], anchos={"Cuenta": 16})
    for celda in wb["Resumen"]["D"][1:]:
        celda.number_format = "0.0%"
    _formatear(wb["Pendientes"], columnas_moneda={"Ingreso", "Egreso"}, columnas_fecha={"Fecha"},
               anchos={"Concepto": 55, "Sugerencia": 70, "Tipo": 34, "Estado": 28})
    for fila in wb["Pendientes"].iter_rows(min_row=2):
        for celda in fila:
            celda.fill = PatternFill("solid", fgColor=FONDO_PENDIENTE)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    _formatear(wb["Por cuenta contable"], anchos={"Cuenta contable": 32},
               columnas_moneda={"Ingreso Banco", "Ingreso Contabilidad", "Dif. ingreso",
                                "Egreso Banco", "Egreso Contabilidad", "Dif. egreso"})
    if "Control de matches" in wb.sheetnames:
        ws = wb["Control de matches"]
        _formatear(ws, anchos={"Nivel": 26},
                   columnas_moneda={"Banco Debe", "Banco Haber", "Contabilidad Debe", "Contabilidad Haber", "Diferencia"})
        col_cierra = [c.value for c in ws[1]].index("Cierra")
        for fila in ws.iter_rows(min_row=2):
            if not fila[col_cierra].value:
                for celda in fila:
                    celda.fill = PatternFill("solid", fgColor=ROJO_ALERTA)
                    celda.font = Font(bold=True)
    for r in resultados:
        ws = wb[r.cuenta[:31]]
        _formatear(ws, columnas_moneda={"Ingreso", "Egreso"}, columnas_fecha={"Fecha", "Fecha registro"},
                   anchos={"Concepto": 60, "Estado": 28, "Nivel": 22, "Cuenta contable": 28, "Contraparte": 28})
        col_conc = [c.value for c in ws[1]].index("Conciliado")
        for fila in ws.iter_rows(min_row=2):
            if fila[col_conc].value:
                for celda in fila:
                    celda.font = Font(color=GRIS_TEXTO)
                    celda.fill = PatternFill("solid", fgColor=FONDO_CONCILIADO)
    wb.save(ruta)
    return ruta
